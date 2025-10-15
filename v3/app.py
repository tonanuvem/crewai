"""
Interface Streamlit para Sistema Multi-Agente BDD
Domain Storytelling → User Stories → BDD → Código

Instalar:
pip install streamlit crewai crewai-tools python-dotenv python-docx pypdf

Executar:
streamlit run app_streamlit.py
"""

import streamlit as st
import os
import time
import logging
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv

# Leitura de arquivos
from docx import Document
from pypdf import PdfReader

# CrewAI
from crewai import Agent, Task, Crew, Process, LLM

# Carregar variáveis de ambiente
load_dotenv()

# Configurar logging
log_dir = Path("logs")
log_dir.mkdir(exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(log_dir / 'crew_logs.log'),
        logging.StreamHandler()
    ],
    force=True
)
logger = logging.getLogger(__name__)

# =============================================================================
# CONFIGURAÇÃO DA PÁGINA
# =============================================================================

st.set_page_config(
    page_title="BDD Multi-Agente | AI Development",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

# CSS Customizado
st.markdown("""
<style>
    .main-header {
        font-size: 2.5rem;
        font-weight: bold;
        color: #1E88E5;
        text-align: center;
        margin-bottom: 1rem;
    }
    .sub-header {
        font-size: 1.2rem;
        color: #666;
        text-align: center;
        margin-bottom: 2rem;
    }
    .agent-card {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 1rem;
        border-radius: 10px;
        color: white;
        margin: 0.5rem 0;
    }
    .success-box {
        background-color: #d4edda;
        border: 1px solid #c3e6cb;
        border-radius: 5px;
        padding: 1rem;
        margin: 1rem 0;
    }
    .error-box {
        background-color: #f8d7da;
        border: 1px solid #f5c6cb;
        border-radius: 5px;
        padding: 1rem;
        margin: 1rem 0;
    }
    .stProgress > div > div > div > div {
        background: linear-gradient(to right, #667eea, #764ba2);
    }
</style>
""", unsafe_allow_html=True)

# =============================================================================
# FUNÇÕES DE LEITURA DE ARQUIVOS
# =============================================================================

def read_word_file(file) -> str:
    """Lê arquivo Word (.docx)"""
    try:
        doc = Document(file)
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        return "\n\n".join(text)
    except Exception as e:
        st.error(f"Erro ao ler arquivo Word: {str(e)}")
        return None


def read_pdf_file(file) -> str:
    """Lê arquivo PDF"""
    try:
        pdf_reader = PdfReader(file)
        text = []
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text.append(page_text)
        return "\n\n".join(text)
    except Exception as e:
        st.error(f"Erro ao ler arquivo PDF: {str(e)}")
        return None


def read_text_file(file) -> str:
    """Lê arquivo de texto (.txt)"""
    try:
        # Garante que o ponteiro de leitura esteja no início
        file.seek(0)
        raw = file.read()
        # uploaded_file.read() retorna bytes em streamlit, então decode quando necessário
        if isinstance(raw, bytes):
            return raw.decode('utf-8', errors='ignore')
        return str(raw)
    except Exception as e:
        st.error(f"Erro ao ler arquivo de texto: {str(e)}")
        return None


def extract_text_from_file(uploaded_file) -> str:
    """Extrai texto de diferentes tipos de arquivo"""
    if uploaded_file is None:
        return None
    
    file_extension = uploaded_file.name.split('.')[-1].lower()
    
    # Rebobinar o arquivo para a função de leitura
    uploaded_file.seek(0)
    
    if file_extension == 'docx':
        return read_word_file(uploaded_file)
    elif file_extension == 'pdf':
        return read_pdf_file(uploaded_file)
    elif file_extension == 'txt':
        return read_text_file(uploaded_file)
    else:
        st.error(f"Formato de arquivo não suportado: .{file_extension}")
        return None


def save_agent_outputs(individual_results: dict, timestamp: datetime):
    """Salva os resultados de cada agente em arquivos específicos"""
    output_dir = Path("outputs") / timestamp.strftime("%Y%m%d_%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    
    file_mappings = {
        "Analista de Domínio Especialista": ("domain_storytelling.md", "md"),
        "Product Owner Ágil": ("user_stories.md", "md"),
        "QA Engineer BDD": ("scenarios.feature", "feature"),
        "Senior Developer": ("implementation.py", "py"),
        "Code Reviewer": ("review_report.md", "md")
    }
    
    saved_files = []
    
    for role, result in individual_results.items():
        if role in file_mappings:
            filename, ext = file_mappings[role]
            file_path = output_dir / filename
            
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(result)
                saved_files.append(str(file_path))
                logger.info(f"Arquivo salvo: {file_path}")
            except Exception as e:
                logger.error(f"Erro ao salvar {file_path}: {e}")
    
    return saved_files, str(output_dir)

# =============================================================================
# CONFIGURAÇÃO DO LLM (compatível CrewAI / LiteLLM)
# =============================================================================

@st.cache_resource
def get_llm(model_choice: str, custom_model: str, temperature: float, api_key: str):
    """
    Cria instância do LLM nativo do CrewAI.
    """
    if not api_key:
        st.error("⚠️ API Key não configurada!")
        st.stop()
    
    # Se o usuário forneceu um modelo custom exato, respeita isso
    if custom_model and custom_model.strip():
        model_name_to_use = custom_model.strip()
    else:
        model_name_to_use = model_choice
    
    # Garantir que o modelo tenha o prefixo gemini/ se for um modelo Gemini
    if not model_name_to_use.startswith("gemini/") and "gemini" in model_name_to_use.lower():
        model_name_to_use = f"gemini/{model_name_to_use}"
    
    try:
        # Cria LLM nativo do CrewAI (LiteLLM)
        llm = LLM(
            model=model_name_to_use,
            api_key=api_key,
            temperature=temperature
        )
        return llm
    except Exception as e:
        # Mensagem amigável para o usuário
        st.error(f"Erro ao inicializar o modelo '{model_name_to_use}': {e}")
        st.info("Confira se o nome do modelo está correto para sua conta Google Vertex/GenAI.")
        st.stop()

# =============================================================================
# CRIAÇÃO DOS AGENTES (mantendo sua estrutura original)
# =============================================================================

def create_agents(llm):
    """Cria os 5 agentes especializados"""
    
    domain_analyst = Agent(
        role="Analista de Domínio Especialista",
        goal="Compreender processos de negócio e criar Domain Storytelling detalhado",
        backstory="""Expert em DDD e Domain Storytelling. Identifica atores, 
        atividades, objetos e sequências em processos de negócio.""",
        llm=llm,
        verbose=True,
        allow_delegation=False
    )
    
    product_owner = Agent(
        role="Product Owner Ágil",
        goal="Transformar Domain Storytelling em User Stories acionáveis",
        backstory="""PO experiente que cria User Stories no formato: 
        Como/Quero/Para com critérios de aceitação claros.""",
        llm=llm,
        verbose=True,
        allow_delegation=False
    )
    
    qa_engineer = Agent(
        role="QA Engineer BDD",
        goal="Criar cenários Gherkin executáveis e completos",
        backstory="""Especialista em BDD que escreve cenários Gherkin 
        cobrindo casos de sucesso, falha e edge cases.""",
        llm=llm,
        verbose=True,
        allow_delegation=False
    )
    
    senior_developer = Agent(
        role="Senior Developer",
        goal="Implementar código limpo baseado em especificações BDD",
        backstory="""Dev sênior expert em Clean Architecture, DDD, TDD e SOLID. 
        Cria código testável e bem documentado.""",
        llm=llm,
        verbose=True,
        allow_delegation=False
    )
    
    code_reviewer = Agent(
        role="Code Reviewer",
        goal="Revisar qualidade do código e sugerir melhorias",
        backstory="""Arquiteto sênior que revisa código, identifica bugs 
        e sugere otimizações seguindo best practices.""",
        llm=llm,
        verbose=True,
        allow_delegation=False
    )
    
    return [domain_analyst, product_owner, qa_engineer, senior_developer, code_reviewer]

# =============================================================================
# CRIAÇÃO DAS TASKS (mantendo sua sequência)
# =============================================================================

def create_tasks(agents, business_description: str):
    """Cria as tasks encadeadas"""
    
    domain_analyst, product_owner, qa_engineer, senior_developer, code_reviewer = agents
    
    task1 = Task(
        description=f"""
        Analise esta descrição de negócio e crie Domain Storytelling:
        
        {business_description}
        
        Produza:
        1. Lista de Atores
        2. Fluxo sequencial (Ator → ação → Objeto)
        3. Entidades do domínio
        4. Regras de negócio
        5. Glossário de termos
        """,
        agent=domain_analyst,
        expected_output="Domain Storytelling completo em Markdown"
    )
    
    task2 = Task(
        description="""
        Transforme o Domain Storytelling em User Stories:
        
        Para cada fluxo:
        - Título
        - Como/Quero/Para
        - Critérios de Aceitação (mínimo 3)
        - Definição de Pronto
        """,
        agent=product_owner,
        expected_output="User Stories em Markdown",
        context=[task1]
    )
    
    task3 = Task(
        description="""
        Crie cenários Gherkin para as User Stories:
        
        Para cada story, crie:
        - Cenário de sucesso
        - Cenário de falha
        - Cenário de edge case
        
        Formato .feature válido
        """,
        agent=qa_engineer,
        expected_output="Cenários Gherkin completos",
        context=[task2]
    )
    
    task4 = Task(
        description="""
        Implemente código Python completo:
        
        Crie:
        1. Entidades (dataclasses)
        2. Use Cases
        3. Step Definitions
        4. Testes unitários
        
        Com type hints e docstrings
        """,
        agent=senior_developer,
        expected_output="Código Python completo",
        context=[task3]
    )
    
    task5 = Task(
        description="""
        Revise o código e produza relatório:
        
        Analise:
        - Aderência às specs
        - Qualidade (SOLID, Clean Code)
        - Cobertura de testes
        - Melhorias sugeridas
        """,
        agent=code_reviewer,
        expected_output="Relatório de revisão",
        context=[task4]
    )
    
    return [task1, task2, task3, task4, task5]

# =============================================================================
# INTERFACE STREAMLIT (restaurada + melhorias LLM)
# =============================================================================

def main():
    # Header
    st.markdown('<p class="main-header">🤖 Sistema Multi-Agente BDD</p>', unsafe_allow_html=True)
    st.markdown(
        '<p class="sub-header">Domain Storytelling → User Stories → Gherkin → Código</p>', 
        unsafe_allow_html=True
    )
    
    # Sidebar - Configurações
    with st.sidebar:
        
        # Informação 
        st.info("💡 Protótipo criado para as disciplinas do Prof. André Sampaio.")
        
        st.divider()
        
        st.header("⚙️ Configurações")
        
        # Campo API Key
        env_api_key = os.getenv("GOOGLE_API_KEY", "")
        api_key = st.text_input(
            "🔑 Google API Key",
            value=env_api_key,
            type="password",
            help="Sua chave da API do Google Gemini"
        )
        
        # Verificar API Key
        if api_key:
            st.success("✅ API Key configurada")
        else:
            st.error("❌ API Key não encontrada")
        
        st.divider()
        
        # Seleção de modelo (mantém opções e permite custom)
        custom_model = st.text_input(
            "Modelo custom (opcional)",
            value="gemini/gemini-2.5-flash",
            help="Exemplos: gemini/gemini-2.5-flash-lite ou gemini/gemini-2.5-pro"
        )
        
        temperature = st.slider(
            "🌡️ Temperatura",
            min_value=0.0,
            max_value=1.0,
            value=0.7,
            step=0.1,
            help="0 = determinístico | 1 = criativo"
        )
        
        st.divider()
        
        # Status dos agentes
        st.header("👥 Agentes")
        agents_info = [
            ("🔍", "Domain Analyst", "Análise de domínio"),
            ("📋", "Product Owner", "User Stories"),
            ("🧪", "QA Engineer", "BDD/Gherkin"),
            ("💻", "Developer", "Código"),
            ("🔎", "Reviewer", "Revisão")
        ]
        
        for icon, name, desc in agents_info:
            st.markdown(f"{icon} **{name}**")
            st.caption(desc)

    
    # Área principal
    tab1, tab2, tab3 = st.tabs(["📄 Input", "🚀 Execução", "📊 Resultados"])
    
    # TAB 1: INPUT
    with tab1:
        st.header("📥 Descrição do Negócio")
        
        # Opção de upload de arquivo
        col1, col2 = st.columns([2, 1])
        
        uploaded_file = None
        with col1:
            uploaded_file = st.file_uploader(
                "Upload de Documento",
                type=['pdf', 'docx', 'txt'],
                help="Formatos: PDF, Word (.docx), Texto (.txt)"
            )
        
        with col2:
            if uploaded_file:
                st.success(f"✅ {uploaded_file.name}")
                # Leitura do tamanho
                try:
                    uploaded_file.seek(0, os.SEEK_END)
                    file_size = uploaded_file.tell() / 1024
                    uploaded_file.seek(0)
                    st.caption(f"Tamanho: {file_size:.1f} KB")
                except:
                    st.caption("Tamanho: N/A")
        
        # Área de texto
        business_description = ""
        if uploaded_file:
            with st.spinner("📖 Extraindo texto do arquivo..."):
                extracted_text = extract_text_from_file(uploaded_file)
                
                if extracted_text:
                    st.success(f"✅ Texto extraído! {len(extracted_text)} caracteres")
                    business_description = st.text_area(
                        "Texto Extraído (editável)",
                        value=extracted_text,
                        height=300,
                        key="extracted_text"
                    )
                else:
                    business_description = ""
        else:
            business_description = st.text_area(
                "Ou digite diretamente aqui:",
                height=300,
                placeholder="""Exemplo:

Nossa empresa precisa de um sistema de aprovação de despesas.

Processo atual:
- Funcionário preenche formulário com valor, categoria e justificativa
- Despesas até R$ 500 são aprovadas automaticamente
- Despesas entre R$ 500 e R$ 5.000 precisam de aprovação do gestor
- Despesas acima de R$ 5.000 precisam de aprovação do diretor
- Sistema deve enviar email para aprovador
- Comprovante é obrigatório para despesas > R$ 100

Regras:
- Não pode haver despesas duplicadas
- Aprovação deve acontecer em até 48h úteis
                """,
                key="manual_text"
            )
        
        # Preview
        if business_description and len(business_description.strip()) > 50:
            st.divider()
            with st.expander("👁️ Preview do Texto", expanded=False):
                st.markdown(business_description)
                
                # Estatísticas
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Caracteres", len(business_description))
                with col2:
                    st.metric("Palavras", len(business_description.split()))
                with col3:
                    st.metric("Linhas", len(business_description.split('\n')))
    
    # TAB 2: EXECUÇÃO
    with tab2:
        st.header("🚀 Executar Processamento")
        
        if not business_description or len(business_description.strip()) < 50:
            st.warning("⚠️ Por favor, forneça uma descrição de negócio na aba Input (mínimo 50 caracteres)")
        elif not api_key:
            st.warning("⚠️ Por favor, configure a API Key na sidebar")
        else:
            st.info(f"📝 Descrição carregada: {len(business_description)} caracteres")
            
            col1, col2 = st.columns([1, 3])
            with col1:
                start_button = st.button(
                    "▶️ Iniciar Processamento",
                    type="primary",
                    use_container_width=True
                )
            
            if start_button:
                # Inicializar session state
                if 'results' not in st.session_state:
                    st.session_state.results = {}
                
                # Container para progresso
                progress_container = st.container()
                
                with progress_container:
                    st.divider()
                    
                    # Progress bar
                    progress_bar = st.progress(0)
                    status_text = st.empty()
                    
                    # Criar placeholders para cada etapa
                    st.subheader("📊 Progresso por Agente")
                    
                    agent_progress = {}
                    try:
                        # Criar LLM
                        status_text.text("🧠 Inicializando modelo Gemini...")
                        llm = get_llm("gemini/gemini-2.5-flash", custom_model, temperature, api_key)
                        progress_bar.progress(10)
                        time.sleep(0.3)
                        
                        # Criar agentes
                        status_text.text("👥 Criando agentes especializados...")
                        agents = create_agents(llm)
                        progress_bar.progress(20)
                        time.sleep(0.3)
                        
                        # Criar tasks
                        status_text.text("📋 Configurando tarefas...")
                        tasks = create_tasks(agents, business_description)
                        progress_bar.progress(30)
                        time.sleep(0.3)
                        
                        # Criar crew
                        status_text.text("🤝 Montando equipe...")
                        crew = Crew(
                            agents=agents,
                            tasks=tasks,
                            process=Process.sequential,
                            verbose=True,
                            memory=True
                        )
                        progress_bar.progress(40)
                        time.sleep(0.3)
                        
                        # Placeholder de Progresso Visual
                        agent_containers = {}
                        for i, (agent, task) in enumerate(zip(agents, tasks)):
                            with st.expander(f"{i+1}. {agent.role}", expanded=True):
                                agent_containers[agent.role] = {
                                    'status': st.empty(),
                                    'result': st.empty()
                                }
                                agent_containers[agent.role]['status'].info(f"⏳ Aguardando...")

                        # Executar
                        status_text.text("⚙️ Executando processamento...")
                        
                        # KICKOFF REAL
                        start_time = time.time()
                        
                        # Inicializar resultados individuais
                        individual_results = {}
                        
                        # Executar tasks sequencialmente e capturar resultados
                        logger.info("Iniciando processamento multi-agente")
                        
                        for i, (agent, task) in enumerate(zip(agents, tasks)):
                            logger.info(f"Executando task {i+1}: {agent.role}")
                            agent_containers[agent.role]['status'].warning(f"🔄 Executando...")
                            
                            # Criar crew temporário para esta task
                            temp_crew = Crew(
                                agents=[agent],
                                tasks=[task],
                                process=Process.sequential,
                                verbose=True
                            )
                            
                            try:
                                task_result = temp_crew.kickoff()
                                result_text = str(task_result)
                                individual_results[agent.role] = result_text
                                
                                # Log do resultado
                                logger.info(f"Task {i+1} ({agent.role}) concluída. Resultado: {result_text[:200]}...")
                                
                                # Atualizar UI
                                agent_containers[agent.role]['status'].success(f"✅ Concluído!")
                                agent_containers[agent.role]['result'].markdown(f"**Resultado:**\n\n{result_text}")
                                
                                progress_bar.progress(40 + (i + 1) * 12)
                                
                            except Exception as e:
                                error_msg = f"Erro na task {i+1}: {str(e)}"
                                logger.error(error_msg)
                                agent_containers[agent.role]['status'].error(f"❌ Erro: {str(e)}")
                                individual_results[agent.role] = f"ERRO: {str(e)}"
                        
                        # Resultado final combinado
                        final_result = "\n\n".join([f"## {role}\n\n{result}" for role, result in individual_results.items()])
                        
                        duration = time.time() - start_time
                        logger.info(f"Processamento concluído em {duration:.1f}s")
                        
                        # Salvar arquivos de saída
                        saved_files, output_dir = save_agent_outputs(individual_results, datetime.now())
                        logger.info(f"Arquivos salvos em: {output_dir}")
                            
                        progress_bar.progress(100)
                        status_text.text("✅ Processamento concluído!")
                        
                        # Salvar resultados
                        st.session_state.results = {
                            'output': final_result,
                            'individual_results': individual_results,
                            'saved_files': saved_files,
                            'output_dir': output_dir,
                            'duration': duration,
                            'timestamp': datetime.now(),
                            'model': custom_model.strip() if custom_model.strip() else "gemini/gemini-2.5-flash",
                            'temperature': temperature
                        }
                        
                        # Mensagem de sucesso
                        st.success(f"🎉 Processamento concluído em {duration:.1f}s!")
                        st.balloons()
                        
                    except Exception as e:
                        st.error(f"❌ Erro durante processamento: {str(e)}")
                        st.exception(e)
    
    # TAB 3: RESULTADOS
    with tab3:
        st.header("📊 Resultados")
        
        if 'results' not in st.session_state or not st.session_state.results:
            st.info("ℹ️ Execute o processamento na aba 'Execução' para ver os resultados")
        else:
            results = st.session_state.results
            
            # Métricas
            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("⏱️ Duração", f"{results['duration']:.1f}s")
            with col2:
                st.metric("🧠 Modelo", results['model'])
            with col3:
                st.metric("🌡️ Temperatura", results['temperature'])
            with col4:
                st.metric("🕐 Horário", results['timestamp'].strftime("%H:%M"))
            
            st.divider()
            
            # Output completo
            with st.expander("📄 Output Completo", expanded=True):
                st.markdown(results['output'])
            
            # Resultados individuais por agente
            if 'individual_results' in results:
                st.subheader("🔍 Resultados por Agente")
                for role, result in results['individual_results'].items():
                    with st.expander(f"📋 {role}", expanded=False):
                        st.markdown(result)
            
            # Arquivos salvos
            if 'saved_files' in results and results['saved_files']:
                st.subheader("📁 Arquivos Gerados")
                st.info(f"📂 Diretório: {results['output_dir']}")
                
                for file_path in results['saved_files']:
                    file_name = Path(file_path).name
                    col1, col2 = st.columns([3, 1])
                    with col1:
                        st.text(f"📄 {file_name}")
                    with col2:
                        try:
                            with open(file_path, 'r', encoding='utf-8') as f:
                                file_content = f.read()
                            st.download_button(
                                "⬇️",
                                data=file_content,
                                file_name=file_name,
                                key=f"download_{file_name}"
                            )
                        except:
                            st.text("❌")
            
            st.divider()
            
            # Opções de download
            st.subheader("💾 Downloads")
            
            col1, col2, col3 = st.columns(3)
            
            with col1:
                st.download_button(
                    label="📥 Download Markdown",
                    data=results['output'],
                    file_name=f"resultado_{results['timestamp'].strftime('%Y%m%d_%H%M%S')}.md",
                    mime="text/markdown"
                )
            
            with col2:
                # Criar versão texto
                text_version = f"""
RESULTADO DO PROCESSAMENTO
Data: {results['timestamp'].strftime('%d/%m/%Y %H:%M:%S')}
Modelo: {results['model']}
Temperatura: {results['temperature']}
Duração: {results['duration']:.1f}s

{results['output']}
                """
                st.download_button(
                    label="📥 Download TXT",
                    data=text_version,
                    file_name=f"resultado_{results['timestamp'].strftime('%Y%m%d_%H%M%S')}.txt",
                    mime="text/plain"
                )
            
            with col3:
                if st.button("🗑️ Limpar Resultados"):
                    st.session_state.results = {}
                    st.rerun()


if __name__ == "__main__":
    main()
